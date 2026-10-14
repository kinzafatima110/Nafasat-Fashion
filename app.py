#!/usr/bin/env python3
"""
NAFASAT FASHION & RENTALS - Digital Showroom & Admin Web Service
Backend application handling public catalogue, admin authentication,
outfit additions with mannequin/dummy photo uploads, and Render hosting.
"""

import os
import json
import re
import shutil
import tempfile
from functools import wraps
from flask import Flask, request, jsonify, session, send_from_directory, redirect
from werkzeug.utils import secure_filename

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_FILE = os.path.join(BASE_DIR, 'data', 'products.json')
ASSETS_DIR = os.path.join(BASE_DIR, 'assets')
IMAGES_DIR = os.path.join(ASSETS_DIR, 'images')

ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'webp', 'gif'}

# Category definitions
CATEGORIES = [
    {"id": "rental-bridal", "name": "Bridal Wear"},
    {"id": "female-partywear", "name": "Party Wear"},
    {"id": "jewelry", "name": "Jewelry"},
    {"id": "kids-rental", "name": "Girls Rental"},
    {"id": "kids-coatpant", "name": "Boys Coat-Pants"},
    {"id": "saree", "name": "Sarees"},
    {"id": "purses", "name": "Clutches & Purses"},
    {"id": "shoes", "name": "Footwear"},
    {"id": "kids-sherwani", "name": "Boys Sherwani"}
]

CATEGORY_MAP = {cat["id"]: cat["name"] for cat in CATEGORIES}

STANDARD_COLORS = [
    "Red & Maroon",
    "Pink & Peach",
    "Blue & Ferozi",
    "Green & Emerald",
    "Gold & Yellow",
    "White & Silver",
    "Black",
    "Purple & Plum",
    "Multi & Other"
]

STANDARD_SIZES = [
    "XS",
    "Small",
    "Medium",
    "Large",
    "XL",
    "Free Size",
    "Kids"
]

app = Flask(__name__, static_folder='.', static_url_path='')
app.secret_key = os.environ.get('SECRET_KEY', 'nafasat_digital_showroom_secret_key_2026')
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50 MB max upload

ADMIN_USERNAME = os.environ.get('ADMIN_USERNAME', 'admin')
ADMIN_PASSWORD = os.environ.get('ADMIN_PASSWORD', 'nafasat2026!')


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def standardize_color(raw_color):
    if not raw_color:
        return "Multi & Other"
    c = raw_color.lower().strip()
    if any(k in c for k in ['red', 'maroon', 'ruby', 'crimson', 'rust', 'burgundy', 'wine']):
        return "Red & Maroon"
    elif any(k in c for k in ['pink', 'peach', 'blush', 'rose', 'coral', 'magenta', 'fuchsia']):
        return "Pink & Peach"
    elif any(k in c for k in ['blue', 'ferozi', 'navy', 'sky', 'teal', 'royal']):
        return "Blue & Ferozi"
    elif any(k in c for k in ['green', 'dhani', 'bottle', 'sea green', 'mint', 'emerald', 'olive']):
        return "Green & Emerald"
    elif any(k in c for k in ['yellow', 'gold', 'mustard', 'lemon', 'fone', 'khaki', 'tan', 'beige']):
        return "Gold & Yellow"
    elif any(k in c for k in ['white', 'silver', 'ivory', 'off white', 'skin', 'zinc', 'cream', 'grey', 'gray']):
        return "White & Silver"
    elif 'black' in c:
        return "Black"
    elif any(k in c for k in ['purple', 'plum', 'violet', 'lavender', 'lilac']):
        return "Purple & Plum"
    else:
        return "Multi & Other"


def standardize_size(raw_size):
    if not raw_size:
        return "Free Size"
    s = raw_size.lower().strip()
    if s in ['small', 's', 'sm'] or 'small' in s:
        return "Small"
    elif s in ['medium', 'm', 'med'] or 'med' in s:
        return "Medium"
    elif s in ['large', 'l', 'lg'] or 'large' in s:
        return "Large"
    elif s in ['xl', 'extra large', 'xxl']:
        return "XL"
    elif s in ['xs', 'extra small']:
        return "XS"
    elif 'kid' in s or re.search(r'\b(24|26|28|30|32|34)\b', s):
        return "Kids"
    elif re.search(r'\b(36)\b', s):
        return "Small"
    elif re.search(r'\b(38)\b', s):
        return "Medium"
    elif re.search(r'\b(40|42)\b', s):
        return "Large"
    else:
        return "Free Size"


def load_products_data():
    if not os.path.exists(DATA_FILE):
        return {"categories": CATEGORIES, "products": []}
    try:
        with open(DATA_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception as e:
        app.logger.error(f"Error loading {DATA_FILE}: {e}")
        return {"categories": CATEGORIES, "products": []}


def save_products_data(data):
    # Recalculate category counts
    cat_counts = {}
    for p in data.get("products", []):
        cid = p.get("category_id")
        cat_counts[cid] = cat_counts.get(cid, 0) + 1

    for cat in data.get("categories", []):
        cat["count"] = cat_counts.get(cat["id"], 0)

    # Atomic write
    os.makedirs(os.path.dirname(DATA_FILE), exist_ok=True)
    tmp_file = f"{DATA_FILE}.tmp.{os.getpid()}"
    with open(tmp_file, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    os.replace(tmp_file, DATA_FILE)


def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('admin_logged_in'):
            return jsonify({'success': False, 'error': 'Unauthorized. Please log in.'}), 401
        return f(*args, **kwargs)
    return decorated_function


# =========================================================================
# PUBLIC WEB ROUTES
# =========================================================================
@app.route('/')
def home():
    return send_from_directory(BASE_DIR, 'index.html')


@app.route('/admin')
def admin_page():
    return send_from_directory(BASE_DIR, 'admin.html')


@app.route('/healthz')
def healthz():
    return jsonify({'status': 'healthy', 'service': 'nafasat-catalogue'})


# =========================================================================
# ADMIN AUTHENTICATION ENDPOINTS
# =========================================================================
@app.route('/api/admin/login', methods=['POST'])
def admin_login():
    data = request.get_json(silent=True) or request.form
    username = (data.get('username') or '').strip()
    password = (data.get('password') or '').strip()

    if username == ADMIN_USERNAME and password == ADMIN_PASSWORD:
        session['admin_logged_in'] = True
        session['admin_username'] = username
        return jsonify({'success': True, 'message': 'Welcome to Nafasat Admin Portal'})
    else:
        return jsonify({'success': False, 'error': 'Invalid username or password'}), 401


@app.route('/api/admin/logout', methods=['POST'])
def admin_logout():
    session.pop('admin_logged_in', None)
    session.pop('admin_username', None)
    return jsonify({'success': True, 'message': 'Logged out successfully'})


@app.route('/api/admin/status', methods=['GET'])
def admin_status():
    logged_in = bool(session.get('admin_logged_in'))
    data = load_products_data()
    return jsonify({
        'logged_in': logged_in,
        'username': session.get('admin_username', '') if logged_in else '',
        'total_items': len(data.get('products', [])),
        'categories': data.get('categories', []),
        'standard_colors': STANDARD_COLORS,
        'standard_sizes': STANDARD_SIZES
    })


# =========================================================================
# ADMIN CATALOG MANAGEMENT ENDPOINTS
# =========================================================================
@app.route('/api/admin/products', methods=['GET'])
@admin_required
def admin_list_products():
    data = load_products_data()
    products = data.get('products', [])
    category = request.args.get('category')
    search = request.args.get('search', '').lower().strip()

    if category and category != 'all':
        products = [p for p in products if p.get('category_id') == category]

    if search:
        products = [
            p for p in products
            if search in p.get('sku', '').lower()
            or search in p.get('color', '').lower()
            or search in p.get('description', '').lower()
            or search in p.get('title', '').lower()
        ]

    return jsonify({
        'total': len(products),
        'products': products
    })


@app.route('/api/admin/add-item', methods=['POST'])
@admin_required
def admin_add_item():
    try:
        sku = (request.form.get('sku') or '').strip().upper()
        category_id = (request.form.get('category_id') or '').strip()
        color = (request.form.get('color') or '').strip()
        user_std_color = (request.form.get('standard_color') or '').strip()
        size = (request.form.get('size') or '').strip()
        user_std_size = (request.form.get('standard_size') or '').strip()
        rent = (request.form.get('rent') or '').strip().replace(',', '').replace('Rs.', '').strip()
        deposit = (request.form.get('deposit') or '').strip().replace(',', '').replace('Rs.', '').strip()
        description = (request.form.get('description') or '').strip()

        if not sku:
            return jsonify({'success': False, 'error': 'Item Code (SKU) is required'}), 400

        if not category_id or category_id not in CATEGORY_MAP:
            return jsonify({'success': False, 'error': 'Valid category is required'}), 400

        category_name = CATEGORY_MAP[category_id]

        # Check for dummy_image
        if 'dummy_image' not in request.files:
            return jsonify({'success': False, 'error': 'Primary Mannequin/Dummy photo is required'}), 400

        dummy_file = request.files['dummy_image']
        if dummy_file.filename == '' or not allowed_file(dummy_file.filename):
            return jsonify({'success': False, 'error': 'Valid dummy image file (jpg/png/webp) is required'}), 400

        # Standardize color and size
        std_color = user_std_color if user_std_color in STANDARD_COLORS else standardize_color(color)
        std_size = user_std_size if user_std_size in STANDARD_SIZES else standardize_size(size)

        # Prepare directory for images
        target_dir = os.path.join(IMAGES_DIR, category_id)
        os.makedirs(target_dir, exist_ok=True)

        # Clean SKU for filename
        safe_sku = re.sub(r'[^A-Za-z0-9_-]', '', sku)
        if not safe_sku:
            safe_sku = f"ITEM_{os.urandom(4).hex()}"

        # Save dummy (primary) image
        dummy_ext = dummy_file.filename.rsplit('.', 1)[1].lower()
        dummy_filename = f"{safe_sku} (2).{dummy_ext}"
        dummy_path = os.path.join(target_dir, dummy_filename)
        dummy_file.save(dummy_path)
        rel_primary_image = f"assets/images/{category_id}/{dummy_filename}"

        saved_images = [rel_primary_image]

        # Save additional angle images
        angle_files = request.files.getlist('angle_images')
        angle_counter = 1
        for f in angle_files:
            if f and f.filename != '' and allowed_file(f.filename):
                if angle_counter == 2:
                    angle_counter += 1  # 2 is reserved for dummy
                ext = f.filename.rsplit('.', 1)[1].lower()
                angle_filename = f"{safe_sku} ({angle_counter}).{ext}"
                angle_path = os.path.join(target_dir, angle_filename)
                f.save(angle_path)
                saved_images.append(f"assets/images/{category_id}/{angle_filename}")
                angle_counter += 1

        # Load existing catalogue data
        data = load_products_data()
        products = data.get('products', [])

        # Check if SKU already exists
        existing_idx = next((i for i, p in enumerate(products) if p.get('sku', '').upper() == sku), None)
        
        # Calculate new product ID
        next_id = 1
        if products:
            next_id = max((p.get('id', 0) for p in products), default=0) + 1

        new_product = {
            "id": next_id if existing_idx is None else products[existing_idx].get('id', next_id),
            "sku": sku,
            "title": f"{category_name} {sku}",
            "category_id": category_id,
            "category_name": category_name,
            "color": color or std_color,
            "standard_color": std_color,
            "size": size or std_size,
            "standard_size": std_size,
            "rent": rent,
            "deposit": deposit,
            "description": description or f"Exclusive {category_name} available for rent.",
            "primary_image": rel_primary_image,
            "images": saved_images,
            "image_count": len(saved_images)
        }

        if existing_idx is not None:
            # Update existing
            products[existing_idx] = new_product
            action = "updated"
        else:
            # Prepend to catalogue so it appears first
            products.insert(0, new_product)
            action = "added"

        data["products"] = products
        save_products_data(data)

        return jsonify({
            'success': True,
            'action': action,
            'message': f"Item {sku} successfully {action} to {category_name}!",
            'product': new_product
        })

    except Exception as e:
        app.logger.exception(e)
        return jsonify({'success': False, 'error': f"Failed to save item: {str(e)}"}), 500


@app.route('/api/admin/delete-item', methods=['POST'])
@admin_required
def admin_delete_item():
    try:
        data_req = request.get_json(silent=True) or request.form
        sku = (data_req.get('sku') or '').strip().upper()
        if not sku:
            return jsonify({'success': False, 'error': 'SKU is required to delete item'}), 400

        catalog = load_products_data()
        products = catalog.get('products', [])

        target_item = next((p for p in products if p.get('sku', '').upper() == sku), None)
        if not target_item:
            return jsonify({'success': False, 'error': f"Item {sku} not found"}), 404

        # Remove item from catalog
        catalog['products'] = [p for p in products if p.get('sku', '').upper() != sku]
        save_products_data(catalog)

        return jsonify({
            'success': True,
            'message': f"Item {sku} deleted successfully from catalogue."
        })

    except Exception as e:
        app.logger.exception(e)
        return jsonify({'success': False, 'error': f"Failed to delete item: {str(e)}"}), 500



# =========================================================================
# EXCEL BULK CATALOGUE SYNC & TEMPLATE DOWNLOAD
# =========================================================================

EXACT_EXCEL_COLUMNS = [
    'SKU',
    'Category',
    'Title',
    'Color',
    'Size',
    'Rent',
    'Deposit',
    'Description',
    'Primary Image',
    'Additional Images'
]

def resolve_category(raw_val):
    if not raw_val:
        return None, None
    clean = str(raw_val).strip().lower()
    for cat in CATEGORIES:
        if clean == cat['name'].lower() or clean == cat['id'].lower():
            return cat['id'], cat['name']
    if 'bridal' in clean:
        return 'rental-bridal', 'Bridal Wear'
    if 'party' in clean:
        return 'female-partywear', 'Party Wear'
    if 'saree' in clean or 'sari' in clean:
        return 'saree', 'Sarees'
    if 'jewel' in clean:
        return 'jewelry', 'Jewelry'
    if 'girl' in clean or ('kid' in clean and 'girl' in clean):
        return 'kids-rental', 'Girls Rental'
    if 'coat' in clean or 'pant' in clean or 'suit' in clean or 'tuxedo' in clean:
        return 'kids-coatpant', 'Boys Coat-Pants'
    if 'purse' in clean or 'clutch' in clean or 'bag' in clean:
        return 'purses', 'Clutches & Purses'
    if 'shoe' in clean or 'heel' in clean or 'footwear' in clean:
        return 'shoes', 'Footwear'
    if 'sherwani' in clean:
        return 'kids-sherwani', 'Boys Sherwani'
    return None, None


@app.route('/api/admin/download-template', methods=['GET'])
def download_excel_template():
    template_path = os.path.join(BASE_DIR, 'data', 'Nafasat_Catalogue_Template.xlsx')
    if not os.path.exists(template_path):
        return jsonify({'success': False, 'error': 'Template file not found.'}), 404
    return send_from_directory(
        os.path.join(BASE_DIR, 'data'),
        'Nafasat_Catalogue_Template.xlsx',
        as_attachment=True,
        download_name='Nafasat_Catalogue_Template.xlsx'
    )


def parse_and_validate_excel_workbook(file_storage):
    """
    Strictly validates columns of an uploaded Excel file.
    Returns (is_valid, error_payload, parsed_rows_or_none)
    """
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_storage, data_only=True)
    except Exception as e:
        return False, {'error': f'Failed to read Excel workbook: {str(e)}. Please ensure the file is a valid .xlsx spreadsheet.'}, None

    # Find sheet: prefer 'Catalogue_Data', otherwise first worksheet
    sheet_names = wb.sheetnames
    target_sheet = None
    for name in sheet_names:
        if name.strip().lower() == 'catalogue_data':
            target_sheet = wb[name]
            break
    if target_sheet is None:
        target_sheet = wb.worksheets[0]

    # Read headers from row 1
    header_cells = [c.value for c in target_sheet[1]]
    raw_headers = [str(h).strip() if h is not None else '' for h in header_cells]
    while raw_headers and not raw_headers[-1]:
        raw_headers.pop()

    norm_expected = [h.strip().upper() for h in EXACT_EXCEL_COLUMNS]
    norm_actual = [h.strip().upper() for h in raw_headers]

    expected_set = set(norm_expected)
    actual_set = set(norm_actual)

    missing_cols = [orig for orig, norm in zip(EXACT_EXCEL_COLUMNS, norm_expected) if norm not in actual_set]
    unexpected_cols = [orig for orig, norm in zip(raw_headers, norm_actual) if norm not in expected_set]

    # Strict Check: exactly the allowed columns, no extra, no missing
    if missing_cols or unexpected_cols or norm_actual != norm_expected:
        error_msg = 'Spreadsheet column format rejected. The uploaded file must contain ONLY the 10 exact accepted columns. Any additional random column or different column names will not be accepted.'
        return False, {
            'error': error_msg,
            'missing_columns': missing_cols,
            'unexpected_columns': unexpected_cols,
            'expected_columns': EXACT_EXCEL_COLUMNS,
            'found_columns': raw_headers
        }, None

    # Parse data rows
    col_indices = {norm_actual[i]: i for i in range(len(norm_actual)) if norm_actual[i] in expected_set}
    rows = list(target_sheet.iter_rows(min_row=2, values_only=True))

    parsed_items = []
    row_errors = []

    for row_num, row_values in enumerate(rows, start=2):
        row_list = list(row_values)
        if not any(cell is not None and str(cell).strip() != '' for cell in row_list):
            continue

        def get_val(col_name):
            idx = col_indices.get(col_name.upper())
            if idx is not None and idx < len(row_list):
                v = row_list[idx]
                return '' if v is None else str(v).strip()
            return ''

        sku = get_val('SKU').upper()
        cat_raw = get_val('Category')
        title = get_val('Title')
        color = get_val('Color')
        size = get_val('Size')
        rent_val = re.sub(r'[^\d]', '', get_val('Rent'))
        deposit_val = re.sub(r'[^\d]', '', get_val('Deposit'))
        desc = get_val('Description')
        primary_img = get_val('Primary Image')
        add_imgs_raw = get_val('Additional Images')

        if not sku:
            row_errors.append(f'Row {row_num}: Missing SKU / Item Code.')
            continue

        cat_id, cat_name = resolve_category(cat_raw)
        if not cat_id:
            row_errors.append(f"Row {row_num} (SKU: {sku}): Unrecognized Category '{cat_raw}'. Allowed categories are: Bridal Wear, Party Wear, Jewelry, Girls Rental, Boys Coat-Pants, Sarees, Clutches & Purses, Footwear, Boys Sherwani.")
            continue

        if not title:
            title = f"{cat_name} {sku}"

        std_color = standardize_color(color)
        std_size = standardize_size(size)

        images_list = []
        if primary_img:
            images_list.append(primary_img)

        if add_imgs_raw:
            for part in re.split(r'[,;]', add_imgs_raw):
                p = part.strip()
                if p and p not in images_list:
                    images_list.append(p)

        if not images_list:
            images_list = ['assets/logo.png']
            primary_img = 'assets/logo.png'

        parsed_items.append({
            'row_num': row_num,
            'sku': sku,
            'category_id': cat_id,
            'category_name': cat_name,
            'title': title,
            'color': color or std_color,
            'standard_color': std_color,
            'size': size or std_size,
            'standard_size': std_size,
            'rent': rent_val,
            'deposit': deposit_val,
            'description': desc or f"Exclusive {cat_name} ensemble available for rent at Nafasat.",
            'primary_image': primary_img,
            'images': images_list,
            'image_count': len(images_list)
        })

    return True, {'row_errors': row_errors}, parsed_items


@app.route('/api/admin/validate-excel', methods=['POST'])
@admin_required
def admin_validate_excel():
    if 'excel_file' not in request.files:
        return jsonify({'success': False, 'error': 'No Excel file was uploaded.'}), 400

    file = request.files['excel_file']
    if not file.filename.lower().endswith(('.xlsx', '.xls')):
        return jsonify({'success': False, 'error': 'Only Excel files (.xlsx, .xls) are accepted.'}), 400

    is_valid, err_or_info, parsed_items = parse_and_validate_excel_workbook(file)
    if not is_valid:
        return jsonify({
            'success': False,
            'error': err_or_info.get('error'),
            'missing_columns': err_or_info.get('missing_columns', []),
            'unexpected_columns': err_or_info.get('unexpected_columns', []),
            'expected_columns': err_or_info.get('expected_columns', []),
            'found_columns': err_or_info.get('found_columns', [])
        }), 400

    row_errors = err_or_info.get('row_errors', [])
    preview = parsed_items[:8] if parsed_items else []

    return jsonify({
        'success': True,
        'message': 'Column format verified successfully!',
        'total_valid_rows': len(parsed_items),
        'row_errors': row_errors,
        'preview': preview
    })


@app.route('/api/admin/upload-excel', methods=['POST'])
@admin_required
def admin_upload_excel():
    try:
        if 'excel_file' not in request.files:
            return jsonify({'success': False, 'error': 'No Excel file was uploaded.'}), 400

        file = request.files['excel_file']
        mode = (request.form.get('mode') or 'merge').strip().lower()

        if not file.filename.lower().endswith(('.xlsx', '.xls')):
            return jsonify({'success': False, 'error': 'Only Excel files (.xlsx, .xls) are accepted.'}), 400

        is_valid, err_or_info, parsed_items = parse_and_validate_excel_workbook(file)
        if not is_valid:
            return jsonify({
                'success': False,
                'error': err_or_info.get('error'),
                'missing_columns': err_or_info.get('missing_columns', []),
                'unexpected_columns': err_or_info.get('unexpected_columns', []),
                'expected_columns': err_or_info.get('expected_columns', []),
                'found_columns': err_or_info.get('found_columns', [])
            }), 400

        if not parsed_items:
            return jsonify({'success': False, 'error': 'The Excel file does not contain any valid data rows to import.'}), 400

        data = load_products_data()
        existing_products = data.get('products', [])

        added_count = 0
        updated_count = 0

        if mode == 'replace':
            new_products_list = []
            for idx, item in enumerate(parsed_items, 1):
                item_copy = dict(item)
                item_copy.pop('row_num', None)
                item_copy['id'] = idx
                new_products_list.append(item_copy)
            data['products'] = new_products_list
            added_count = len(new_products_list)
        else:
            sku_map = {p.get('sku', '').upper(): i for i, p in enumerate(existing_products) if p.get('sku')}

            for item in parsed_items:
                item_copy = dict(item)
                item_copy.pop('row_num', None)
                sku = item_copy['sku']

                if sku in sku_map:
                    target_idx = sku_map[sku]
                    item_copy['id'] = existing_products[target_idx].get('id', target_idx + 1)
                    existing_products[target_idx] = item_copy
                    updated_count += 1
                else:
                    item_copy['id'] = len(existing_products) + 1
                    existing_products.insert(0, item_copy)
                    added_count += 1
                    sku_map[sku] = 0

            for idx, p in enumerate(existing_products, 1):
                p['id'] = idx
            data['products'] = existing_products

        cat_counts = {}
        for p in data['products']:
            cid = p.get('category_id')
            if cid:
                cat_counts[cid] = cat_counts.get(cid, 0) + 1

        for c in data.get('categories', []):
            c['count'] = cat_counts.get(c['id'], 0)

        save_products_data(data)

        return jsonify({
            'success': True,
            'message': f"Excel sync complete! {added_count} outfits added, {updated_count} updated. Total catalogue size: {len(data['products'])} items.",
            'added_count': added_count,
            'updated_count': updated_count,
            'total_items': len(data['products'])
        })

    except Exception as e:
        app.logger.exception(e)
        return jsonify({'success': False, 'error': f"Failed to sync Excel: {str(e)}"}), 500


# Fallback for static assets in development
@app.route('/<path:filename>')
def serve_static(filename):
    return send_from_directory(BASE_DIR, filename)


if __name__ == '__main__':
    port = int(os.environ.get('PORT', 8080))
    print(f"================================================================")
    print(f" Nafasat Fashion & Rentals - Digital Showroom & Admin Portal")
    print(f" Running at: http://localhost:{port}")
    print(f" Admin portal: http://localhost:{port}/admin")
    print(f" Default Admin: username '{ADMIN_USERNAME}', password '{ADMIN_PASSWORD}'")
    print(f"================================================================")
    app.run(host='0.0.0.0', port=port, debug=False)
